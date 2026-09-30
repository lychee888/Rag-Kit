"""Real files/storage with deterministic models; required in CI."""
import json
from types import SimpleNamespace

import numpy as np
import pytest
from docx import Document
from PIL import Image
from typer.testing import CliRunner

from rag_kit.cli import main as cli
from rag_kit.config import Config
from rag_kit.ingest import pipeline
from rag_kit.store import VectorStore


@pytest.mark.parametrize('mixed', [False, True], ids=['image-only', 'mixed'])
@pytest.mark.parametrize('surface', ['cli', 'watcher', 'polling'])
def test_caption_failure_preserves_whole_source(tmp_path, monkeypatch, capsys, mixed, surface):
    image = tmp_path / 'chart.png'
    Image.new('RGB', (100, 100), 'blue').save(image)
    path = tmp_path / 'chart.docx'
    doc = Document()
    if mixed:
        doc.add_paragraph('Original document text')
    doc.add_picture(str(image))
    doc.save(path)
    cfg = Config.from_dict({'db_path': str(tmp_path / 'db'), 'watch_folder': str(tmp_path), 'vlm_enabled': True})
    monkeypatch.setattr(cli, '_ensure_config', lambda: cfg)
    monkeypatch.setattr(cli, 'create_engine_from_config', lambda _: SimpleNamespace(
        embed_texts=lambda texts: np.tile([[1., 0., 0., 0.]], (len(texts), 1))))
    monkeypatch.setattr(pipeline, '_run_vlm_captioning', lambda *a: [{
        'text': 'Revenue chart', 'page': 0, 'vlm_generated': True, 'source_type': 'image'}])
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(cfg.db_path)
    before = store._get_table().to_arrow().to_pylist()
    assert any(row['vlm_generated'] for row in before)
    if mixed:
        doc.paragraphs[0].text = 'Changed text must not replace old snapshot'
        doc.save(path)

    def fail(*args):
        raise RuntimeError('injected captioning exception')
    monkeypatch.setattr(pipeline, '_run_vlm_captioning', fail)
    if surface == 'cli':
        response = runner.invoke(cli.app, ['ingest', str(path), '--json'])
        assert response.exit_code == 1
        body = json.loads(response.stdout)
    elif surface == 'watcher':
        cli._WatchHandler(cfg, json_output=True)._ingest_one(str(path))
        body = json.loads(capsys.readouterr().out)
        assert body['event'] == 'ingest_failed'
        assert body['source_replaced'] is False
    else:
        from rag_kit import watcher
        monkeypatch.setattr('rag_kit.config.get_config', lambda: cfg)
        monkeypatch.setattr(watcher, '_load_state', lambda: {str(path): 'old-hash'})
        saved = []
        monkeypatch.setattr(watcher, '_save_state', lambda value: saved.append(dict(value)))
        cfg.supported_extensions = ['.docx']
        body = watcher.scan_and_ingest(once=True)
        assert saved == [{str(path): 'old-hash'}]
    assert body['status'] == 'partial'
    assert body['errors']
    assert store._get_table().to_arrow().to_pylist() == before


@pytest.mark.parametrize('source_filter', [False, True])
def test_keyword_only_never_constructs_embedding_engine(tmp_path, monkeypatch, source_filter):
    cfg = Config.from_dict({'db_path': str(tmp_path / 'db'), 'search_alpha': 0})
    source = str(tmp_path / 'doc.txt')
    VectorStore(cfg.db_path).replace_source(source, [{'id': 'a', 'text': 'needle', 'source': source}], np.array([[1., 0.]]))
    monkeypatch.setattr(cli, '_ensure_config', lambda: cfg)
    def forbidden(*args):
        raise AssertionError('keyword search must not construct an embedding engine')
    monkeypatch.setattr(cli, 'create_engine_from_config', forbidden)
    args = ['query', 'needle', '--json'] + (['--source', source] if source_filter else [])
    result = CliRunner().invoke(cli.app, args)
    assert result.exit_code == 0, result.exception
    assert json.loads(result.stdout)['results'][0]['text'] == 'needle'


@pytest.mark.parametrize('available', [False, True])
def test_unavailable_or_incomplete_captions_raise(tmp_path, monkeypatch, available):
    import rag_kit.vlm as vlm
    import rag_kit.vlm.extractor as extractor
    monkeypatch.setattr(extractor, 'extract_images_from_docx', lambda path: [object()])
    monkeypatch.setattr(vlm, 'VLMCaptioner', lambda: SimpleNamespace(
        is_available=lambda: available, caption_regions=lambda *a, **kw: []))
    with pytest.raises(pipeline.IncompleteExtractionError):
        pipeline._run_vlm_captioning(str(tmp_path / 'visual.docx'), '.docx', ['en'])

"""Tests for new rag-kit capabilities: xlsx/pptx ingestion, recursive
folder ingest, hybrid search, and delete-by-full-path.
"""

from __future__ import annotations

import numpy as np

from rag_kit.ingest.pipeline import ingest_file, ingest_folder, SUPPORTED_EXTENSIONS
from rag_kit.store import VectorStore


# ── Supported extension registry ──────────────────────────────────────


def test_supported_extensions_include_new_formats():
    assert ".xlsx" in SUPPORTED_EXTENSIONS
    assert ".pptx" in SUPPORTED_EXTENSIONS
    assert ".png" in SUPPORTED_EXTENSIONS
    assert ".jpg" in SUPPORTED_EXTENSIONS


# ── XLSX ingestion ───────────────────────────────────────────────────


def test_ingest_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "value"])
    ws.append(["widget", 42])
    ws.append(["gadget", 7.5])
    p = tmp_path / "sheet.xlsx"
    wb.save(str(p))

    chunks = ingest_file(str(p), use_ocr=False, use_vlm=False)
    assert chunks, "xlsx should produce chunks"
    text = " ".join(c["text"] for c in chunks)
    assert "widget" in text
    assert "gadget" in text
    assert all(c["source"].endswith("sheet.xlsx") for c in chunks)


# ── PPTX ingestion ──────────────────────────────────────────────────


def test_ingest_pptx(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only layout
    slide.shapes.title.text = "Roadmap"
    body = slide.shapes.add_textbox(0, 0, 100, 100)
    body.text_frame.text = "Q3 release notes here"
    p = tmp_path / "deck.pptx"
    prs.save(str(p))

    chunks = ingest_file(str(p), use_ocr=False, use_vlm=False)
    assert chunks, "pptx should produce chunks"
    text = " ".join(c["text"] for c in chunks)
    assert "Roadmap" in text
    assert "release notes" in text


# ── Recursive folder ingest ─────────────────────────────────────────


def test_ingest_folder_is_recursive(tmp_path):
    (tmp_path / "docs" / "sub" / "deep").mkdir(parents=True)
    (tmp_path / "docs" / "readme.md").write_text("top level doc", encoding="utf-8")
    (tmp_path / "docs" / "sub" / "nested.txt").write_text("nested doc", encoding="utf-8")
    (tmp_path / "docs" / "sub" / "deep" / "deepest.md").write_text(
        "deepest doc", encoding="utf-8")

    chunks = ingest_folder(str(tmp_path / "docs"), use_ocr=False, use_vlm=False)
    texts = " ".join(c["text"] for c in chunks)
    assert "top level doc" in texts
    assert "nested doc" in texts
    assert "deepest doc" in texts


# ── Hybrid search (FTS + vector) ───────────────────────────────────


def test_hybrid_search_finds_keyword_match(tmp_path):
    store = VectorStore(db_path=str(tmp_path / "db"))
    dim = 8
    chunks = [
        {"id": "a1", "text": "quantum entanglement experiment", "source": "/tmp/a.pdf", "page": 1, "chunk_idx": 0},
        {"id": "b1", "text": "recipe for chocolate cake", "source": "/tmp/b.pdf", "page": 1, "chunk_idx": 0},
    ]
    vectors = np.random.rand(2, dim).astype(np.float32)
    store.add_chunks(chunks, vectors)

    # Keyword-only query that the random vectors won't rank highly: the
    # hybrid path should still surface the cake recipe via FTS.
    query_vec = np.zeros(dim, dtype=np.float32)
    results = store.search_hybrid("chocolate cake recipe", query_vec, top_k=2, alpha=0.0)
    scores = {r["id"]: r["score"] for r in results}
    assert "b1" in scores, "FTS should surface exact keyword match"

    # Vector-only (alpha=1) still works as plain vector search.
    results_v = store.search_hybrid("x", query_vec, top_k=2, alpha=1.0)
    assert len(results_v) == 2


# ── Delete by full path ───────────────────────────────────────────


def test_delete_matches_stored_full_path(tmp_path):
    store = VectorStore(db_path=str(tmp_path / "db"))
    chunks = [
        {"id": "a1", "text": "hello", "source": str(tmp_path / "a.txt"), "page": 1, "chunk_idx": 0},
        {"id": "b1", "text": "world", "source": str(tmp_path / "b.txt"), "page": 1, "chunk_idx": 0},
    ]
    vectors = np.random.rand(2, 8).astype(np.float32)
    store.add_chunks(chunks, vectors)

    removed = store.delete_by_source(str(tmp_path / "a.txt"))
    assert removed == 1
    assert store.count_rows() == 1

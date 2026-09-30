"""Exercise failures below captioning, with real documents and source storage."""
import io
import json
from types import SimpleNamespace

import fitz
import numpy as np
import pytest
from PIL import Image
from typer.testing import CliRunner

from rag_kit.cli import main as cli
from rag_kit.config import Config
from rag_kit.ingest import pipeline
from rag_kit.store import VectorStore
from rag_kit.vlm import extractor


def make_pdf(path, kind='vector'):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), 'Native text with enough characters to avoid OCR.')
    if kind in ('vector', 'mixed'):
        for i in range(5):
            page.draw_rect(fitz.Rect(50 + i * 45, 100, 80 + i * 45, 200 + i * 20))
    if kind in ('raster', 'mixed', 'two-images'):
        for i, color in enumerate(['blue', 'red'] if kind == 'two-images' else ['blue']):
            buf = io.BytesIO()
            Image.new('RGB', (100, 100), color).save(buf, format='PNG')
            page.insert_image(fitz.Rect(50 + i * 220, 350, 250 + i * 220, 550), stream=buf.getvalue())
    doc.save(path)
    doc.close()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    cfg = Config.from_dict({'db_path': str(tmp_path / 'db'), 'watch_folder': str(tmp_path), 'vlm_enabled': True})
    monkeypatch.setattr(cli, '_ensure_config', lambda: cfg)
    monkeypatch.setattr(cli, 'create_engine_from_config', lambda _: SimpleNamespace(
        embed_texts=lambda texts: np.tile([[1., 0., 0., 0.]], (len(texts), 1))))
    captioner = SimpleNamespace(is_available=lambda: True, caption_regions=lambda regions, **kw: [
        {'text': 'Visual caption', 'page': r.page_num, 'vlm_generated': True, 'source_type': 'image'} for r in regions])
    monkeypatch.setattr('rag_kit.vlm.VLMCaptioner', lambda: captioner)
    return cfg


@pytest.mark.parametrize('surface', ['cli', 'watcher', 'polling'])
@pytest.mark.parametrize('kind', ['vector', 'mixed'])
def test_render_failure_preserves_whole_snapshot(setup, tmp_path, monkeypatch, capsys, surface, kind):
    path = tmp_path / 'chart.pdf'
    make_pdf(path, kind)
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(setup.db_path)
    before = store._get_table().to_arrow().to_pylist()
    assert any(r['vlm_generated'] for r in before)
    monkeypatch.setattr(fitz.Page, 'get_pixmap', lambda *a, **kw: (_ for _ in ()).throw(MemoryError('render failed')))
    if surface == 'cli':
        response = runner.invoke(cli.app, ['ingest', str(path), '--json'])
        assert response.exit_code == 1
        body = json.loads(response.stdout)
    elif surface == 'watcher':
        cli._WatchHandler(setup, json_output=True)._ingest_one(str(path))
        body = json.loads(capsys.readouterr().out)
        assert body['source_replaced'] is False
    else:
        from rag_kit import watcher
        monkeypatch.setattr('rag_kit.config.get_config', lambda: setup)
        monkeypatch.setattr(watcher, '_load_state', lambda: {str(path): 'old-hash'})
        saved = []
        monkeypatch.setattr(watcher, '_save_state', lambda value: saved.append(dict(value)))
        setup.supported_extensions = ['.pdf']
        body = watcher.scan_and_ingest(once=True)
        assert saved == [{str(path): 'old-hash'}]
    assert body['status'] == 'partial'
    assert body['errors']
    assert store._get_table().to_arrow().to_pylist() == before


def test_complete_raster_fallback_succeeds_when_page_render_fails(setup, tmp_path, monkeypatch):
    path = tmp_path / 'raster.pdf'
    make_pdf(path, 'raster')
    monkeypatch.setattr(fitz.Page, 'get_pixmap', lambda *a, **kw: (_ for _ in ()).throw(MemoryError('render failed')))
    with fitz.open(path) as doc:
        regions = extractor.extract_images_from_pdf(doc)
    assert len(regions) == 1
    assert regions[0].source_type == 'embedded_image'
    result = CliRunner().invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 0, result.exception
    assert any(r['vlm_generated'] for r in VectorStore(setup.db_path)._get_table().to_arrow().to_pylist())


def test_partial_embedded_image_fallback_does_not_replace_source(setup, tmp_path, monkeypatch):
    path = tmp_path / 'images.pdf'
    make_pdf(path, 'two-images')
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(setup.db_path)
    before = store._get_table().to_arrow().to_pylist()
    with fitz.open(path) as doc:
        second_xref = doc[0].get_images(full=True)[1][0]
    original = fitz.Document.extract_image
    def extract(doc, xref):
        if xref == second_xref:
            raise MemoryError('second image failed')
        return original(doc, xref)
    monkeypatch.setattr(fitz.Page, 'get_pixmap', lambda *a, **kw: (_ for _ in ()).throw(MemoryError('render failed')))
    monkeypatch.setattr(fitz.Document, 'extract_image', extract)
    result = runner.invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 1
    assert json.loads(result.stdout)['status'] == 'partial'
    assert store._get_table().to_arrow().to_pylist() == before


@pytest.mark.parametrize('kind', ['blank', 'short-text', 'separator'])
@pytest.mark.parametrize('use_ocr', [False, True])
def test_blank_and_native_text_pages_need_no_ocr_or_vlm(tmp_path, monkeypatch, kind, use_ocr):
    path = tmp_path / 'document.pdf'
    doc = fitz.open()
    page = doc.new_page()
    if kind == 'short-text':
        page.insert_text((50, 50), 'Title')
    elif kind == 'separator':
        page.insert_text((50, 50), 'Native text with enough characters to avoid OCR.')
        doc.new_page()
    doc.save(path)
    doc.close()
    def forbidden(*a, **kw):
        raise AssertionError('no visual content needs a model')
    monkeypatch.setattr(pipeline, '_ocr_pdf_page', forbidden)
    monkeypatch.setattr(pipeline, '_vlm_pdf_page', forbidden)
    monkeypatch.setattr('rag_kit.vlm.VLMCaptioner', lambda: SimpleNamespace(is_available=lambda: False))
    chunks = pipeline.ingest_file(path, use_ocr=use_ocr, use_vlm=True)
    if kind == 'blank':
        assert chunks == []
    else:
        assert chunks[0]['text'].startswith('Title' if kind == 'short-text' else 'Native text')


def test_successfully_empty_pdf_removes_previous_rows(setup, tmp_path, monkeypatch):
    path = tmp_path / 'empty.pdf'
    doc = fitz.open(); doc.new_page(); doc.save(path); doc.close()
    store = VectorStore(setup.db_path)
    store.replace_source(str(path), [{'id': 'old', 'text': 'Old content', 'source': str(path)}], [[1., 0., 0., 0.]])
    monkeypatch.setattr('rag_kit.vlm.VLMCaptioner', lambda: SimpleNamespace(is_available=lambda: False))
    result = CliRunner().invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 0, result.exception
    assert store.count_rows() == 0


def test_scanned_content_still_fails_with_unavailable_vlm(tmp_path, monkeypatch):
    path = tmp_path / 'scan.pdf'
    doc = fitz.open(); page = doc.new_page()
    buf = io.BytesIO(); Image.new('RGB', (100, 100), 'blue').save(buf, format='PNG')
    page.insert_image(page.rect, stream=buf.getvalue()); doc.save(path); doc.close()
    monkeypatch.setattr(pipeline, '_ocr_pdf_page', lambda *a, **kw: '')
    def unavailable(*a, **kw):
        raise RuntimeError('model unavailable')
    monkeypatch.setattr('rag_kit.vlm.VLMCaptioner', lambda: SimpleNamespace(caption_image=unavailable))
    with pytest.raises(pipeline.IncompleteExtractionError):
        pipeline.ingest_file(path, use_vlm=True)


def test_docx_image_read_failure_is_reported(setup, tmp_path, monkeypatch):
    from docx import Document
    from docx.parts.image import ImagePart
    image = tmp_path / 'chart.png'
    Image.new('RGB', (100, 100), 'blue').save(image)
    path = tmp_path / 'chart.docx'
    doc = Document(); doc.add_paragraph('Native document text'); doc.add_picture(str(image)); doc.save(path)
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(setup.db_path)
    before = store._get_table().to_arrow().to_pylist()
    def fail(part):
        raise MemoryError('image read failed')
    monkeypatch.setattr(ImagePart, 'blob', property(fail))
    result = runner.invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 1
    assert json.loads(result.stdout)['status'] == 'partial'
    assert store._get_table().to_arrow().to_pylist() == before


def test_inline_image_cannot_be_hidden_by_successful_xobject_fallback(setup, tmp_path, monkeypatch):
    path = tmp_path / 'inline.pdf'
    make_pdf(path, 'raster')
    with fitz.open(path) as doc:
        page = doc[0]
        contents = page.get_contents()
        xref = doc.get_new_xref(); doc.update_object(xref, '<<>>')
        doc.update_stream(xref, b'q\n200 0 0 200 300 100 cm\nBI /W 1 /H 1 /BPC 8 /CS /RGB /F /AHx ID\nFF0000>\nEI\nQ\n')
        doc.xref_set_key(page.xref, 'Contents', '[' + ' '.join(f'{ref} 0 R' for ref in [*contents, xref]) + ']')
        doc.save(tmp_path / 'complete.pdf')
    path = tmp_path / 'complete.pdf'
    with fitz.open(path) as doc:
        assert len(doc[0].get_images(full=True)) == 1
        assert len(doc[0].get_image_info()) == 2
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(setup.db_path)
    before = store._get_table().to_arrow().to_pylist()
    monkeypatch.setattr(fitz.Page, 'get_pixmap', lambda *a, **kw: (_ for _ in ()).throw(MemoryError('render failed')))
    result = runner.invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 1
    assert json.loads(result.stdout)['status'] == 'partial'
    assert store._get_table().to_arrow().to_pylist() == before


def test_checked_form_widget_still_requires_visual_extraction(tmp_path, monkeypatch):
    path = tmp_path / 'checkbox.pdf'
    doc = fitz.open(); page = doc.new_page()
    widget = fitz.Widget()
    widget.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
    widget.field_name = 'approved'
    widget.field_value = 'Yes'
    widget.rect = fitz.Rect(50, 50, 80, 80)
    page.add_widget(widget); doc.save(path); doc.close()
    monkeypatch.setattr(pipeline, '_ocr_pdf_page', lambda *a, **kw: '')
    calls = []
    monkeypatch.setattr(pipeline, '_vlm_pdf_page', lambda *a, **kw: calls.append(True) or 'Approved checkbox is checked')
    monkeypatch.setattr('rag_kit.vlm.VLMCaptioner', lambda: SimpleNamespace(is_available=lambda: False))
    chunks = pipeline.ingest_file(path, use_vlm=True)
    assert calls == [True]
    assert chunks[0]['text'] == 'Approved checkbox is checked'
    assert chunks[0]['vlm_generated'] is True


@pytest.mark.parametrize('mask_kind', ['soft', 'explicit'])
def test_masked_chart_cannot_be_replaced_by_base_image(setup, tmp_path, monkeypatch, mask_kind):
    from PIL import ImageDraw
    path = tmp_path / 'alpha-chart.pdf'
    image = (Image.new('RGBA', (200, 200), (0, 0, 0, 0)) if mask_kind == 'soft'
             else Image.new('RGB', (200, 200), 'white'))
    draw = ImageDraw.Draw(image)
    for i in range(4):
        draw.rectangle((10 + i * 45, 180 - (i + 1) * 35, 40 + i * 45, 190), fill='black')
    buf = io.BytesIO(); image.save(buf, format='PNG')
    doc = fitz.open(); page = doc.new_page()
    page.insert_text((50, 50), 'Native text with enough characters to avoid OCR.')
    page.insert_image(fitz.Rect(50, 100, 350, 400), stream=buf.getvalue())
    image_info = page.get_images(full=True)[0]
    if mask_kind == 'soft':
        assert image_info[1] > 0
    else:
        assert image_info[1] == 0
        doc.xref_set_key(image_info[0], 'Mask', '[255 255 255 255 255 255]')
        assert doc.xref_get_key(image_info[0], 'Mask')[0] == 'array'
    doc.save(path); doc.close()
    runner = CliRunner()
    assert runner.invoke(cli.app, ['ingest', str(path), '--json']).exit_code == 0
    store = VectorStore(setup.db_path)
    before = store._get_table().to_arrow().to_pylist()
    monkeypatch.setattr(fitz.Page, 'get_pixmap', lambda *a, **kw: (_ for _ in ()).throw(MemoryError('render failed')))
    result = runner.invoke(cli.app, ['ingest', str(path), '--json'])
    assert result.exit_code == 1
    assert json.loads(result.stdout)['status'] == 'partial'
    assert store._get_table().to_arrow().to_pylist() == before
